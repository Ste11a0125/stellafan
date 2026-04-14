from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# A4 page setup
for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)

# Style helpers
def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
    p.runs[0].font.size = Pt(15)
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = RGBColor(0x2E, 0x86, 0xAB)
    p.runs[0].font.size = Pt(12)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def spacer():
    p = doc.add_paragraph('')
    p.paragraph_format.space_after = Pt(2)

# ── FRONT PAGE ──────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
run = p.add_run('ADOVA')
run.bold = True
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('[Tagline / strapline — elevator pitch in one line]')
run.italic = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x2E, 0x86, 0xAB)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(60)
run = p.add_run('BUSINESS PLAN  |  2025')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Confidential — Prepared by the Adova Founding Team')
run.italic = True
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

doc.add_page_break()

# ── SECTION 2: EXECUTIVE SUMMARY ────────────────────────────
h1('Executive Summary')
h2('The Opportunity in One Sentence')
bullet('[Single sentence capturing problem, solution, and scale]')

h2('The Problem')
bullet('[What structural gap Adova addresses]')
bullet('[Key stat: 200M+ SMEs globally without in-house marketing]')

h2('The Solution')
bullet('[How Adova works — plain language input → AI → human review → delivery]')
bullet('[Key value props: 5× faster, up to 60% cheaper, no learning curve]')

h2('Market Opportunity')
bullet('[TAM / SAM / SOM headline figures]')
bullet('[Primary market: UK/EU SMEs]')

h2('Business Model')
bullet('[Three-tier pricing: Starter, Growth £29/mo, Scale custom]')
bullet('[Trajectory: Service → Product → Platform]')

h2('Vision and Mission')
bullet('Vision: [one sentence]')
bullet('Mission: [one sentence]')

h2('Next Steps')
bullet('[What Adova is seeking — pilots, partnerships, investment]')

doc.add_page_break()

# ── SECTION 3: PROBLEM DEFINITION ───────────────────────────
h1('Idea Generation and Problem Definition')

h2('The Structural Gap in SME Marketing')
bullet('[Context: 200M+ SMEs, most without marketing function]')
bullet('[Why this is a structural problem, not a motivational one]')

h2('Three Converging Barriers')
bullet('1. No in-house marketing capability — [explain]')
bullet('2. Agency costs are prohibitive — [explain with estimated cost range]')
bullet('3. Generic AI tools are too complex — [explain]')

h2('The Gap Adova Occupies')
bullet('[How Adova sits between expensive agencies and inaccessible DIY tools]')
bullet('[Why now — market conditions, AI maturity, e-commerce growth]')

spacer()

# ── SECTION 4: OBJECTIVES ───────────────────────────────────
h1('Objectives')

h2('North Star Metric')
bullet('[Metric name and definition]')
bullet('[Why this metric was chosen over alternatives]')

h2('Supporting KPIs')
bullet('KPI 1: [name and definition]')
bullet('KPI 2: [name and definition]')
bullet('KPI 3: [name and definition]')

h2('12-Month OKRs')
bullet('O1: [Objective] — KR1 / KR2 / KR3')
bullet('O2: [Objective] — KR1 / KR2 / KR3')
bullet('O3: [Objective] — KR1 / KR2 / KR3')
bullet('O4: [Objective] — KR1 / KR2 / KR3')

doc.add_page_break()

# ── SECTION 5: SALES AND MARKETING PLAN ─────────────────────
h1('Sales and Marketing Plan')

h2('Go-to-Market Strategy')
bullet('Phase 1 — Direct Outreach (Months 1–3): [detail]')
bullet('Phase 2 — Demo-Based Acquisition (Months 2–6): [detail]')
bullet('Phase 3 — Pilot Projects (Months 1–6): [detail]')
bullet('Phase 4 — Case Study-Driven Promotion (Months 4–12): [detail]')

h2('Founder-Led Sales Approach')
bullet('[How the founding team runs early sales personally]')
bullet('[When and how a sales function will be added]')

h2('Pricing Justification')
bullet('Starter (pay-per-use): [what it replaces and why priced here]')
bullet('Growth (£29/month, 10 pieces): [what it replaces and why priced here]')
bullet('Scale (custom, unlimited): [what it replaces and why priced here]')
bullet('[Comparison table placeholder — Adova vs agency vs freelancer vs raw AI]')

spacer()

# ── SECTION 6: MARKET ANALYSIS ──────────────────────────────
h1('Market Analysis')

h2('Market Size: TAM / SAM / SOM')
bullet('TAM: [global SME figure and estimated marketing spend]')
bullet('SAM: [UK + EU focus and serviceable segment]')
bullet('SOM: [realistic capture in Years 1–3]')
bullet('Note: all figures are estimates — [label assumptions]')

h2('SWOT Analysis')
bullet('Strengths: [list]')
bullet('Weaknesses: [list]')
bullet('Opportunities: [list]')
bullet('Threats: [list]')

h2('Competitor Analysis')
bullet('[Table placeholder: Adova vs Traditional Agency vs Freelancer vs Generic AI tools]')
bullet('[Highlight the gap in the middle Adova occupies]')

doc.add_page_break()

# ── SECTION 7: STRATEGIC MANAGEMENT PLAN ────────────────────
h1('Strategic Management Plan')

h2('Vision and Mission')
bullet('Vision: [one sentence]')
bullet('Mission: [one sentence]')

h2('Core Competencies')
bullet('1. [Competency name]: [description]')
bullet('2. [Competency name]: [description]')
bullet('3. [Competency name]: [description]')
bullet('4. [Competency name]: [description]')

h2('Desired Outcomes by Phase')
bullet('Phase 1 — Service (Months 1–12): [outcomes]')
bullet('Phase 2 — Product (Months 13–24): [outcomes]')
bullet('Phase 3 — Platform (Months 25–36+): [outcomes]')

h2('Strategic Frameworks')
bullet('Jobs-to-be-Done: [how applied]')
bullet('Lean Validation Loops: [how applied]')
bullet('OKR Framework: [how applied]')
bullet('Service Blueprint: [how applied]')

h2('KPIs by Phase')
bullet('[Table placeholder: KPI × Phase 1 / Phase 2 / Phase 3]')

h2('RACI Matrix')
bullet('[Table placeholder: Activity × Founding Team / AI System / Human Review / Client]')
bullet('Key: R = Responsible | A = Accountable | C = Consulted | I = Informed')

doc.add_page_break()

# ── SECTION 8: FINANCIALS ────────────────────────────────────
h1('Financial Projections (Years 1–5)')

h2('Revenue and Cost Projections')
bullet('[Table placeholder: rows = revenue lines + cost lines + profit; cols = Year 1–5]')
bullet('Revenue lines: Starter, Growth subscriptions, Scale')
bullet('Cost lines: AI infrastructure, human review, platform dev, customer acquisition, operations')
bullet('Bottom line: operating profit / (loss) per year')

h2('Assumption Justification')
bullet('Client growth curve: [basis for Year 1–5 client count assumptions]')
bullet('Scale tier pricing: [assumption — custom pricing estimated at £X/month]')
bullet('Starter average transaction: [assumption — estimated at £X per request]')
bullet('Human review team cost: [assumption — X part-time reviewers at £X/month]')
bullet('AI infrastructure cost: [assumption — scales with usage volume]')
bullet('Customer acquisition cost: [assumption — estimated as X% of revenue]')
bullet('Break-even: [projected at Month X — basis for assumption]')

spacer()

# ── SECTION 9: APPENDIX — BUSINESS MODEL CANVAS ─────────────
h1('Appendix — Business Model Canvas')

h2('Business Model Canvas: Adova')
bullet('[Full BMC table placeholder — 9 blocks]')

blocks = [
    'Key Partners: [AI API providers, e-commerce platforms, human review contractors, trade associations]',
    'Key Activities: [content generation, quality review, client onboarding, platform development, sales]',
    'Key Resources: [AI technology, domain knowledge, human review team, brand reputation]',
    'Value Propositions: [professional marketing at SME prices, no learning curve, 5× faster, human-reviewed quality]',
    'Customer Relationships: [managed service (Phase 1), self-service with support (Phase 3)]',
    'Channels: [direct outreach, LinkedIn, trade associations, referrals, e-commerce platform integrations]',
    'Customer Segments: [small manufacturers, workshop/makers, solo e-commerce sellers, export businesses]',
    'Cost Structure: [AI infrastructure, human review team, platform development, sales and marketing]',
    'Revenue Streams: [pay-per-use fees, monthly subscriptions (£29–custom), enterprise contracts]',
]
for b in blocks:
    bullet(b)

# ── SAVE ─────────────────────────────────────────────────────
output_path = '/home/user/stellafan/Adova_Business_Plan_Skeleton.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
