from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── COLOURS ─────────────────────────────────────────────────
NAVY        = (0x1B, 0x4F, 0x72)
BLUE        = (0x2E, 0x86, 0xAB)
SLATE       = (0x7F, 0x8C, 0x8D)
DARK        = (0x1A, 0x1A, 0x1A)

# ── HELPERS ──────────────────────────────────────────────────
def rgb(t): return RGBColor(*t)

def add_paragraph(doc, text='', bold=False, italic=False,
                  size=10.5, color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT,
                  space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold   = bold
        run.italic = italic
        run.font.size  = Pt(size)
        run.font.name  = 'Calibri'
        run.font.color.rgb = rgb(color)
    return p

def add_run(para, text, bold=False, italic=False,
            size=10.5, color=DARK):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(size)
    run.font.name  = 'Calibri'
    run.font.color.rgb = rgb(color)
    return run

def h1(doc, text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    for run in p.runs:
        run.font.name  = 'Calibri'
        run.font.size  = Pt(15)
        run.font.color.rgb = rgb(NAVY)
        run.font.bold  = True
    return p

def h2(doc, text):
    p = doc.add_heading(text, level=2)
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after  = Pt(4)
    for run in p.runs:
        run.font.name  = 'Calibri'
        run.font.size  = Pt(11.5)
        run.font.color.rgb = rgb(BLUE)
        run.font.bold  = True
    return p

def bullet(doc, text, sub=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if sub:
        p.paragraph_format.left_indent = Inches(0.4)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.name = 'Calibri'
    run.font.color.rgb = rgb(DARK)
    return p

def bullet_kv(doc, key, value, sub=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if sub:
        p.paragraph_format.left_indent = Inches(0.4)
    r1 = p.add_run(key)
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.name = 'Calibri'
    r1.font.color.rgb = rgb(DARK)
    r2 = p.add_run(value)
    r2.font.size = Pt(10.5)
    r2.font.name = 'Calibri'
    r2.font.color.rgb = rgb(DARK)
    return p

def body(doc, text, space_after=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(space_after)
    for run in p.runs:
        run.font.size = Pt(10.5)
        run.font.name = 'Calibri'
        run.font.color.rgb = rgb(DARK)
    return p

def note(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.size  = Pt(9)
        run.font.name  = 'Calibri'
        run.font.color.rgb = rgb(SLATE)
        run.italic = True
    return p

def add_horizontal_rule(doc, color_hex='2E86AB'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def set_cell_bg(cell, hex_str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_str)
    tcPr.append(shd)

def cell_text(cell, text, bold=False, size=9.5,
              color='1A1A1A', align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(size)
    run.font.name  = 'Calibri'
    run.font.color.rgb = RGBColor(*bytes.fromhex(color))

def make_table(doc, rows, cols, col_widths=None):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    if col_widths:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Cm(col_widths[i])
    return table

# ════════════════════════════════════════════════════════════
# BUILD DOCUMENT
# ════════════════════════════════════════════════════════════
doc = Document()

for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width  = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.0)

# ── FRONT PAGE ───────────────────────────────────────────────

# Large vertical spacer before company name
add_paragraph(doc, space_before=60, space_after=0)

# Company name
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(6)
run = p.add_run('ADOVA')
run.bold = True
run.font.size = Pt(40)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(NAVY)

# Divider
add_horizontal_rule(doc)

# Tagline
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after  = Pt(8)
run = p.add_run('Build great products. Let Adova tell their story.')
run.italic = True
run.font.size = Pt(15)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(BLUE)

# Divider
add_horizontal_rule(doc)

# Elevator pitch
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after  = Pt(0)
run = p.add_run(
    'An AI-powered marketing assistant built for small manufacturers,\n'
    'workshops, and independent product sellers — delivering professional-grade\n'
    'content at a fraction of agency cost, with no learning curve required.'
)
run.font.size = Pt(11)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(DARK)

# Spacer before document metadata
add_paragraph(doc, space_before=50, space_after=0)

# Document type
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(6)
run = p.add_run('BUSINESS PLAN')
run.bold = True
run.font.size = Pt(13)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(NAVY)

# Year
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(10)
run = p.add_run('2025')
run.font.size = Pt(11)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(SLATE)

# Confidential
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(4)
run = p.add_run('Confidential — Prepared by the Adova Founding Team')
run.italic = True
run.font.size = Pt(9)
run.font.name = 'Calibri'
run.font.color.rgb = rgb(SLATE)

doc.add_page_break()

# ════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════
out = '/home/user/stellafan/Adova_Business_Plan.docx'
doc.save(out)
print(f'Saved: {out}')
